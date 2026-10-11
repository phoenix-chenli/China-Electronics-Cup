"""Reproduce curated public source fixtures. Never creates model-written facts.

Run with the document workspace runtime, not the portable app startup runtime.
The original PDFs and original ZIP remain unchanged.
"""
from pathlib import Path
from decimal import Decimal
import argparse
import hashlib
import json
import re
import zipfile
import io
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '02_数据资料' / '公开统计分析'


def main(bundle, data_only=False):
    catalog = json.loads((OUT / '来源清单.json').read_text())
    sources = {x['source_id']: x for x in catalog['sources']}
    pages = {}
    for sid, source in sources.items():
        path = OUT / source['file_path']
        assert hashlib.sha256(path.read_bytes()).hexdigest() == source['original_sha256']
        pages[sid] = [p.extract_text() for p in PdfReader(path).pages]
    observations, indicators = [], {}

    def add(sid, region, metric, title, value, unit, page, year=2024, population=None, basis=None):
        value = str(value)
        assert Decimal(value).is_finite()
        quote = re.sub(r'\s+', ' ', pages[sid][page - 1]).strip()
        assert value in quote, (sid, metric, value, page)
        indicators.setdefault(metric, {'metric_id': metric, 'metric_name': title, 'unit': unit,
                                       'definition': title + '，以对应公报说明和观测行口径为准。'})
        observations.append({'region_id': region, 'year': year, 'metric_id': metric,
            'value_decimal': value, 'unit': unit, 'population_type': population,
            'price_basis': basis, 'preliminary': True, 'source_id': sid,
            'source_pdf_page': page, 'source_quote': quote,
            'extract_status': 'curated_native_pdf_text', 'review_status': 'pending_independent_agent_visual_review'})

    def economic(sid, region, values, growth, shares, page=1):
        price = '绝对数按当年价格，增长速度按不变价格' if region != 'shijingshan' else '公报明示增速按不变价格，绝对数未另行说明'
        add(sid, region, 'gdp', '地区生产总值 GDP', values[0], '亿元', page, basis=price)
        add(sid, region, 'gdp_growth', '地区生产总值实际增长率', growth[0], '%', page, basis='不变价格')
        for i, label in enumerate(['第一产业', '第二产业', '第三产业'], 1):
            if values[i] is None:
                continue
            add(sid, region, f'sector{i}_value', label + '增加值', values[i], '亿元', page, basis=price)
            add(sid, region, f'sector{i}_growth', label + '增加值实际增长率', growth[i], '%', page, basis='不变价格')
            add(sid, region, f'sector{i}_share', label + '占 GDP 比重 公布值', shares[i-1], '%', page, basis='公布比例')

    h, x, b = 'huainan_stats_2024', 'xinjin_stats_2024', 'shijingshan_stats_2024'
    economic(h, 'huainan', ['1716.0', '153.0', '681.6', '881.4'], ['5.0', '2.6', '5.4', '5.1'], ['8.9', '39.7', '51.4'])
    economic(x, 'xinjin', ['534.94', '17.46', '224.83', '292.65'], ['5.9', '2.6', '8.7', '4.0'], ['3.3', '42.0', '54.7'])
    economic(b, 'shijingshan', ['1312.9', None, '90.8', '1222.1'], ['6.5', None, '-3.1', '7.3'], [None, '6.9', '93.1'], 2)
    for sid, region, value, rate, page in [(h, 'huainan', '532.3', '3.0', 1), (x, 'xinjin', '154.31', '7.8', 2), (b, 'shijingshan', '30.7', '-7.2', 2)]:
        add(sid, region, 'industrial_value', '工业增加值', value, '亿元', page, basis='以公报绝对数口径为准')
        add(sid, region, 'industrial_growth', '工业增加值实际增长率', rate, '%', page, basis='不变价格')
    for metric, title, value, unit in [
        ('private_value', '民营经济增加值', '354.57', '亿元'),
        ('private_growth', '民营经济增加值实际增长率', '8.3', '%'),
        ('private_share', '民营经济占 GDP 比重 公布值', '66.3', '%'),
        ('private_sector1_value', '第一产业民营经济增加值', '15.94', '亿元'),
        ('private_sector2_value', '第二产业民营经济增加值', '152.62', '亿元'),
        ('private_sector3_value', '第三产业民营经济增加值', '186.01', '亿元'),
        ('private_industrial_value', '民营工业经济增加值', '126.58', '亿元'),
        ('private_industrial_share', '民营工业占全口径工业增加值比重', '82.0', '%')]:
        add(x, 'xinjin', metric, title, value, unit, 1, basis='当年价格' if unit == '亿元' else '公报增长口径' if metric.endswith('growth') else '公布比例')
    add(x, 'xinjin', 'gdp', '地区生产总值 GDP', '503.12', '亿元', 8, year=2023, basis='第五次经济普查后修订绝对数，不能据此冒充实际增速')
    add(x, 'xinjin', 'private_value', '民营经济增加值', '332.89', '亿元', 8, year=2023, basis='第五次经济普查后修订绝对数')
    for metric, title, value, unit, page in [
        ('population_resident', '年末常住人口', '301.90', '万人', 3),
        ('population_urban', '常住城镇人口', '194.54', '万人', 3),
        ('population_rural', '常住乡村人口', '107.36', '万人', 3),
        ('population_male', '常住男性人口', '154.45', '万人', 3),
        ('population_female', '常住女性人口', '147.45', '万人', 3),
        ('population_age60plus', '60周岁及以上常住人口', '75.29', '万人', 3),
        ('urbanization_share', '常住人口城镇化比例 表内精确公布值', '64.44', '%', 3)]:
        add(h, 'huainan', metric, title, value, unit, page, population='常住人口')
    for metric, title, value in [
        ('population_registered', '户籍人口', '32.90'), ('population_urban', '户籍城镇人口', '18.75'),
        ('population_rural', '户籍乡村人口', '14.15'), ('population_male', '户籍男性人口', '16.10'),
        ('population_female', '户籍女性人口', '16.80')]:
        add(x, 'xinjin', metric, title, value, '万人', 7, population='户籍人口')
    add(b, 'shijingshan', 'population_resident', '年末常住人口', '56.1', '万人', 3, population='常住人口')
    add(b, 'shijingshan', 'population_nonlocal', '常住外来人口', '15.4', '万人', 3, population='常住外来人口')
    add(b, 'shijingshan', 'population_nonlocal_share', '常住外来人口占常住人口比重', '27.5', '%', 3, population='常住人口')
    for sid, region, fine, coarse, page in [(h, 'huainan', '40', '65.2', 15), (x, 'xinjin', '31', '49', 7), (b, 'shijingshan', '31.1', '60', 11)]:
        add(sid, region, 'pm25', 'PM2.5 年均浓度', fine, '微克/立方米', page)
        add(sid, region, 'pm10', 'PM10 年均浓度', coarse, '微克/立方米', page)
    for metric, title, value in [('information_value', '信息传输软件和信息技术服务业增加值', '568.5'), ('finance_value', '金融业增加值', '175.8'), ('construction_value', '建筑业增加值', '61.1')]:
        add(b, 'shijingshan', metric, title, value, '亿元', 2)
    for metric, title, value in [('information_share', '信息传输软件和信息技术服务业占GDP比重', '43.3'), ('finance_share', '金融业占GDP比重', '13.4'), ('construction_share', '建筑业占GDP比重', '4.7')]:
        add(b, 'shijingshan', metric, title, value, '%', 2)
    for indicator in indicators.values():
        key=indicator['metric_id']
        if key in {'pm25','pm10'}:
            indicator['metric_name']=indicator['metric_name'].replace('年均浓度','浓度 公报公布值')
            indicator['definition']='公报年度浓度，须核对观测行 statistical_scope；不能补造原句没有明示的年均口径。'
        elif key in {'population_urban','population_rural','population_male','population_female'}:
            indicator['metric_name']=indicator['metric_name'].replace('常住','').replace('户籍','')
            indicator['definition']='必须同时匹配观测行 population_type；户籍人口与常住人口不能直接混合比较。'
    for row in observations:
        if row['metric_id'] in {'pm25','pm10'}:
            row['statistical_scope']='公报年度浓度，原句未显式写年均' if row['region_id']=='xinjin' else '公报明确的年均浓度'
            if row['region_id']=='xinjin':
                row['additional_pdf_pages']=[6]
                row['additional_source_quotes']=[{'source_pdf_page':6,'source_quote':re.sub(r'\s+',' ',pages[x][5]).strip()}]
                row['source_note']='原句跨原PDF第6至7页；第6页列PM2.5，第7页续写PM10及数值，原句没有明示年均。'
        if row['region_id']=='shijingshan' and row['metric_id'] in {'sector2_value','industrial_value','construction_value'}:
            row['source_note']='原表第二产业90.8，工业30.7和建筑业61.1之和91.8，与第二产业差1.0亿元。保留原值，不擅自更正；需核实原公报。'

    (OUT / '指标核验.json').write_text(json.dumps({'schema_version': '1.0',
        'regions': [{'region_id': 'huainan', 'region_name': '淮南市'}, {'region_id': 'xinjin', 'region_name': '新津区'}, {'region_id': 'shijingshan', 'region_name': '石景山区'}],
        'indicators': list(indicators.values()), 'observations': observations,
        'review_boundary': '原件文字提取与代理视觉核验不等于真人复核；human_review=pending'}, ensure_ascii=False, indent=2))

    if data_only:
        print(json.dumps({'observations':len(observations),'data_only':True}))
        return

    derived = OUT / '转换展示件'
    derived.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    for style in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        font = doc.styles[style].font
        font.name, font.color.rgb = 'Songti SC', RGBColor(0, 0, 0)
        doc.styles[style]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')
        for border in list(doc.styles[style]._element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
        font.underline=False
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.add_paragraph('新津区2024年统计公报原文摘录', 'Title')
    doc.add_paragraph('本文件由官方 PDF 的真实原文转换，用于验证 Word 文档解析、段落定位和跨格式来源追踪。机构发布的原件是 PDF；此 Word 是项目制作的转换展示件，未由模型编写统计内容。')
    doc.add_paragraph('来源 成都市新津区人民政府\n原件 SHA256 ' + sources[x]['original_sha256'])
    doc.add_paragraph('官方地址 ' + sources[x]['official_url'])
    for page in [1, 7, 8]:
        doc.add_page_break()
        doc.add_paragraph(f'原 PDF 第{page}页', 'Heading 1')
        text = pages[x][page - 1].strip()
        text = re.sub(r'^—\s*\d+\s*—\s*', '', text)
        # Preserve native paragraph text; line wrapping is presentation only.
        for paragraph in re.split(r'(?<=。)\s*|(?=\[\d\])', text):
            if paragraph.strip():
                doc.add_paragraph(re.sub(r'\s+', ' ', paragraph).strip())
    word_path = derived / '新津区2024年统计公报原文摘录_转换展示件.docx'
    doc.core_properties.author = '证流资料整理'
    doc.core_properties.title = '新津区2024年统计公报原文摘录'
    doc.core_properties.comments = '转换件 原始发布机构与原件页码详见正文 不等于官方Word原件'
    doc.save(word_path)

    scan = catalog['scan_candidate']
    with zipfile.ZipFile(bundle) as archive:
        candidates = [n for n in archive.namelist() if n.endswith('/mee_environment_2024.pdf')]
        data = next(data for name in candidates if hashlib.sha256(data := archive.read(name)).hexdigest() == scan['original_sha256'])
    assert hashlib.sha256(data).hexdigest() == scan['original_sha256']
    mapping = [1, 10, 11, 12]
    original, excerpt = PdfReader(io.BytesIO(data)), PdfWriter()
    for page in mapping:
        excerpt.add_page(original.pages[page - 1])
    excerpt.add_metadata({'/Title': '2024中国生态环境状况公报 原页摘录', '/Author': '原件 中华人民共和国生态环境部', '/Subject': '派生摘页 不是完整公报 原件页码 1 10 11 12'})
    scan_path = derived / '2024中国生态环境状况公报_扫描原页摘录.pdf'
    with scan_path.open('wb') as stream:
        excerpt.write(stream)
    catalog['derived_documents'] = [
        {'source_id': x, 'title': word_path.name, 'file_path': str(word_path.relative_to(OUT)), 'origin': 'faithful_native_text_excerpt_docx',
         'official_url': sources[x]['official_url'], 'institution': sources[x]['institution'], 'original_sha256': sources[x]['original_sha256'],
         'file_sha256': hashlib.sha256(word_path.read_bytes()).hexdigest(), 'original_pdf_pages': [1, 7, 8], 'document_year': 2024,
         'note': '由官方PDF真实原文转换，不是机构发布的Word；Word按逻辑节定位，正文保留原PDF页号'},
        {**{k: scan[k] for k in ['source_id', 'title', 'institution', 'official_url', 'original_sha256']},
         'file_path': str(scan_path.relative_to(OUT)), 'file_sha256': hashlib.sha256(scan_path.read_bytes()).hexdigest(),
         'origin': 'original_scanned_pdf_page_excerpt', 'document_year': 2024, 'original_pdf_pages': mapping,
         'note': '从150页扫描原件摘取4页，页序与原件映射保留，原件没有文字层，需要真实OCR；不把摘页当全文'}]
    (OUT / '来源清单.json').write_text(json.dumps(catalog, ensure_ascii=False, indent=2))
    print(json.dumps({'observations': len(observations), 'indicators': len(indicators), 'derived_documents': 2}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--bundle', required=True)
    parser.add_argument('--data-only', action='store_true')
    args=parser.parse_args()
    main(args.bundle,args.data_only)
