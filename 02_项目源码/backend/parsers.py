"""Bounded native extraction and local OCR, retaining original page/cell evidence."""
from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile
from contextlib import contextmanager
from pathlib import Path
from xml.etree import ElementTree as ET
from . import table_semantics, document_structure, formula_structure

PARSER_VERSION = "zhiheng-multimodal-3.4"
MAX_OCR_PAGES = 40
MAX_DOCUMENT_SECONDS = 180
PDF_PAGE_TIMEOUT = 20
MAX_BYTES = 80 * 1024 * 1024
MAX_PAGES = 500
MAX_TEXT = 2_000_000
MAX_ROWS = 50_000
MAX_COLUMNS = 256
MAX_CELLS = 1_000_000
MAX_ZIP_BYTES = 256 * 1024 * 1024
SUPPORTED = {"pdf", "docx", "txt", "md", "csv", "xlsx", "png", "jpg", "jpeg", "sqlite", "db"}
KINDS = {"csv": "spreadsheet", "xlsx": "spreadsheet", "db": "database", "sqlite": "database", "png": "image", "jpg": "image", "jpeg": "image"}


class ParseError(ValueError):
    pass


def quoted(value):
    return '"' + value.replace('"', '""') + '"'


def value_json(value):
    if isinstance(value, bytes):
        return {"type": "blob", "bytes": len(value), "hex_preview": value[:32].hex()}
    if isinstance(value, str) and len(value) > 12000:
        return value[:12000] + "…（单元格预览已截断）"
    if isinstance(value, float):
        import math
        if not math.isfinite(value):
            return str(value)
    return value


@contextmanager
def readonly_database(path):
    """Input copies never acquire write connections, even when they contain hostile schema."""
    connection = sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro&immutable=1", uri=True, timeout=3)
    deadline = time.monotonic() + 12
    connection.set_progress_handler(lambda: int(time.monotonic() > deadline), 2000)
    connection.execute("PRAGMA query_only=ON")
    connection.execute("PRAGMA trusted_schema=OFF")
    denied = {sqlite3.SQLITE_INSERT, sqlite3.SQLITE_UPDATE, sqlite3.SQLITE_DELETE, sqlite3.SQLITE_ATTACH, sqlite3.SQLITE_DETACH, sqlite3.SQLITE_CREATE_TABLE, sqlite3.SQLITE_DROP_TABLE, sqlite3.SQLITE_ALTER_TABLE, sqlite3.SQLITE_CREATE_TRIGGER, sqlite3.SQLITE_DROP_TRIGGER, sqlite3.SQLITE_CREATE_VIEW, sqlite3.SQLITE_DROP_VIEW}
    connection.set_authorizer(lambda action, a, b, db, trigger: sqlite3.SQLITE_DENY if action in denied or (action == sqlite3.SQLITE_FUNCTION and str(b).lower() in {"load_extension", "writefile", "readfile"}) else sqlite3.SQLITE_OK)
    try:
        yield connection
    finally:
        connection.close()


def observed_type(values):
    values = [value for value in values if value is not None and value != ""]
    if not values:
        return "空列"
    if all(isinstance(value, (int, bool)) for value in values):
        return "整数"
    if all(isinstance(value, (int, float)) for value in values):
        return "小数"
    if all(isinstance(value, str) for value in values):
        if all(re.fullmatch(r"-?(?:0|[1-9]\d*)", value.strip()) for value in values):
            return "整数（文本存储）"
        if all(re.fullmatch(r"-?(?:0|[1-9]\d*)\.\d+", value.strip()) for value in values):
            return "小数（文本存储）"
        if all(re.match(r"^\d{4}-\d{2}-\d{2}(?:T| |$)", value) for value in values):
            return "日期文本"
        return "文本"
    return "混合"


def sqlite_tables(path, observe_types=False):
    with readonly_database(path) as connection:
        if connection.execute("PRAGMA quick_check(1)").fetchone()[0] != "ok":
            raise ParseError("SQLite 完整性检查未通过，请重新导出数据库后导入")
        records = connection.execute("SELECT name,sql FROM sqlite_schema WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name LIMIT 257").fetchall()
        if len(records) > 256:
            raise ParseError("数据库超过 256 张表，请按业务拆分后导入")
        tables, notices = [], []
        for name, declaration in records:
            if declaration and declaration.lstrip().upper().startswith("CREATE VIRTUAL TABLE"):
                notices.append(f"虚拟表「{name}」未加载，仅支持普通数据表")
                continue
            columns = [{"name": r[1], "type": r[2] or "未声明"} for r in connection.execute(f"PRAGMA table_xinfo({quoted(name)})") if r[6] in {0, 2, 3}]
            if len(columns) > MAX_COLUMNS:
                notices.append(f"表「{name}」超过 {MAX_COLUMNS} 列，未加载")
                continue
            row_count = connection.execute(f"SELECT COUNT(*) FROM {quoted(name)}").fetchone()[0]
            rows = connection.execute(f"SELECT * FROM {quoted(name)} LIMIT 200").fetchall()
            if observe_types:
                for index, column in enumerate(columns):
                    column["type"] = observed_type([row[index] for row in rows])
            tables.append(dict(name=name, columns=columns, row_count=row_count, preview=[[value_json(x) for x in row] for row in rows[:5]]))
        if connection.execute("SELECT 1 FROM sqlite_schema WHERE type='view' LIMIT 1").fetchone():
            notices.append("视图未执行；当前只展示数据库中的普通表")
        return tables, notices


def table_page(path, table_name, allowed, offset, limit):
    if table_name not in allowed:
        raise KeyError(table_name)
    with readonly_database(path) as connection:
        actual = connection.execute("SELECT sql FROM sqlite_schema WHERE type='table' AND name=?", (table_name,)).fetchone()
        if not actual or str(actual[0]).lstrip().upper().startswith("CREATE VIRTUAL TABLE"):
            raise KeyError(table_name)
        columns = [{"name": r[1], "type": r[2] or "未声明"} for r in connection.execute(f"PRAGMA table_xinfo({quoted(table_name)})") if r[6] in {0, 2, 3}]
        # Table name is both allowlisted and quoted; pagination is bound, never interpolated.
        from .table_engine import stable_order_for_table
        ordering = stable_order_for_table(connection, table_name)
        rows = connection.execute(f"SELECT * FROM {quoted(table_name)} ORDER BY {ordering} LIMIT ? OFFSET ?", (limit, offset)).fetchall()
        count = connection.execute(f"SELECT COUNT(*) FROM {quoted(table_name)}").fetchone()[0]
        if isinstance(allowed, dict):
            columns = allowed[table_name]
    return dict(name=table_name, columns=columns, rows=[[value_json(x) for x in row] for row in rows], row_count=count, offset=offset, limit=limit)


def check_zip(path):
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
                raise ParseError("文档解压后过大，请拆分或重新导出后导入")
            if any(i.flag_bits & 1 for i in infos):
                raise ParseError("不支持加密压缩文档，请解密后导入")
            if any(i.file_size > 1024 * 1024 and i.file_size / max(i.compress_size, 1) > 300 for i in infos):
                raise ParseError("文档压缩比异常，为避免内存耗尽已停止解析")
    except zipfile.BadZipFile as exc:
        raise ParseError("文件不是有效的 Office 文档，请从原软件重新导出") from exc


def page(number, title, text, image=False):
    return dict(number=number, title=title, text=text, char_count=len(text), has_image=image)


def base_result():
    return dict(pages=[], tables=[], warnings=[], metadata={}, page_count=0, summary="", data_path=None)


def decode_text(data):
    if b"\x00" in data:
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            return data.decode("utf-16"), "UTF-16"
        raise ParseError("文件包含二进制内容，不能按文本解析")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    raise ParseError("无法识别文本编码，请另存为 UTF-8 后重试")


def clean_headers(values):
    result, used = [], set()
    for i, value in enumerate(values):
        candidate = str(value).strip()[:160] if value is not None else ""
        candidate = candidate or f"列{i + 1}"
        base, suffix = candidate, 2
        while candidate.casefold() in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


class TableWriter:
    def __init__(self, path):
        self.path = path
        path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path)
        self.cells = 0
        self.names = set()

    def close(self):
        self.connection.close()

    def add(self, name, headers, rows):
        if not headers:
            return
        if len(headers) > MAX_COLUMNS:
            raise ParseError(f"表格超过 {MAX_COLUMNS} 列，请拆分后导入")
        if name in self.names:
            name += f"_{len(self.names) + 1}"
        self.names.add(name)
        names = clean_headers(headers)
        # Empty declared type preserves native XLSX number/string values; CSV remains text.
        self.connection.execute(f"CREATE TABLE {quoted(name)} ({','.join(quoted(h) for h in names)})")
        statement = f"INSERT INTO {quoted(name)} VALUES({','.join('?' for _ in names)})"
        for count, row in enumerate(rows, start=1):
            if count > MAX_ROWS or self.cells + len(names) > MAX_CELLS:
                raise ParseError(f"表格超过 {MAX_ROWS:,} 行或累计 {MAX_CELLS:,} 单元格，请拆分后导入")
            if len(row) > len(names):
                raise ParseError("表格数据行的列数超过表头，请修正文件后重试")
            converted = []
            for value in list(row) + [None] * (len(names) - len(row)):
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                if isinstance(value, str) and len(value) > 32768:
                    raise ParseError("单元格超过 32,768 字符，请拆分长文本后导入")
                converted.append(value)
            self.connection.execute(statement, converted)
            self.cells += len(names)
        self.connection.commit()


def _pdf_worker(path, index=None, image_path=None, timeout=PDF_PAGE_TIMEOUT, cancelled=lambda: False):
    command = [sys.executable, "-m", "backend.pdf_extract", str(path)]
    if index is not None:
        command.extend([str(index), str(image_path)])
    try:
        from .reliability import run_cancellable
        process = run_cancellable(command, cwd=Path(__file__).resolve().parent.parent,
                                  cancelled=cancelled, timeout=max(.1, timeout))
    except subprocess.TimeoutExpired as exc:
        raise ParseError("PDF 页面解析超时，可拆分页面后重试") from exc
    if len(process.stdout) > 24_000_000:
        raise ParseError("PDF 页面解析结果超过安全上限")
    try:
        # MuPDF may print diagnostics before its JSON result.
        value = json.loads(process.stdout.decode("utf-8").splitlines()[-1])
    except (ValueError, IndexError, UnicodeError) as exc:
        raise ParseError("PDF 无法读取，请检查文件或重新导出后重试") from exc
    if process.returncode or value.get("error"):
        raise ParseError(value.get("error") or "PDF 页面解析失败，可重试")
    return value


def _save_regions(result, regions, destination):
    """PDF/image grids retain every source row; headers are hints, never guessed SQL schema."""
    result["metadata"]["table_regions"] = regions
    _region_semantics(result, regions)
    if not regions:
        return
    writer = TableWriter(destination)
    try:
        for region in regions:
            rows = region.pop("rows")
            writer.add(region["table_name"], [f"列{i + 1}" for i in range(region["column_count"])], rows)
    finally:
        writer.close()
    result["tables"], notices = sqlite_tables(destination, observe_types=True)
    result["data_path"] = str(destination)
    result["warnings"].extend(notices)
    result["metadata"]["table_row_note"] = "PDF / 图片表格保留全部原始行，列名为列1、列2等；表头与合并覆盖空值未转成业务数据，计算前需确认表头、单位和合计行。"
    result["warnings"].append(result["metadata"]["table_row_note"])
    if any(r["extraction"] == "ocr_grid" for r in regions):
        result["warnings"].append("扫描表格根据真实格线定位；单元格文字来自 OCR，数字及空格须对照原件复核")


def _region_semantics(result, regions):
    semantics = []
    pages = {p['number']: p for p in result['pages']}
    for region in regions[:table_semantics.MAX_TABLES]:
        current = pages.get(region['page'], {})
        siblings = [r for r in regions if r['page'] == region['page']]
        before, after = table_semantics.page_context(region, siblings, current)
        semantics.append(table_semantics.build_table_semantics(region,
            format='image' if result['metadata'].get('width') else 'pdf',
            source_position={'page': region['page'], 'page_unit': 'physical_page', 'bbox': region['bbox']},
            before=before, after=after))
    result['metadata']['table_semantics'] = semantics
    result['metadata']['table_semantics_truncated'] = len(regions) > len(semantics)
    if len(regions) > len(semantics):
        result['warnings'].append(f'仅为前{table_semantics.MAX_TABLES}张表保存语义关系，其余表仍保留原始网格')


def _quality_metadata(pages, truncated=False):
    from .ocr_native import CONFIDENCE_NOTE
    return dict(native_pages=sum(p.get("extraction") == "native" for p in pages),
                ocr_pages=sum(p.get("extraction") == "ocr" for p in pages),
                ocr_comparison_pages=sum(bool(p.get("ocr_comparison")) for p in pages),
                unreadable_pages=sum(p.get("extraction") == "unreadable" for p in pages),
                warning_pages=[p["number"] for p in pages if p.get("quality", {}).get("warnings")],
                ocr_confidence_note=CONFIDENCE_NOTE, processing_truncated=truncated,
                missing_page_note="仅检查文件实际包含的页面；无法据此判断纸质原稿是否缺页。")


def parse_pdf(path, progress, destination=None, checkpoint_dir=None, cancelled=lambda: False, *, defer_pages=(), defer_table_write=False):
    from PIL import Image
    from .ocr_native import OCRError, quality_metrics, text_blocks
    from .ocr_adapter import recognize
    from .import_batches import PageCheckpoints
    from . import adaptive_processing as adaptive
    result = base_result()
    started = time.monotonic()
    info = _pdf_worker(path, cancelled=cancelled)
    count = info["page_count"]
    cache = PageCheckpoints(checkpoint_dir or destination.parent / '.page-checkpoints', path, PARSER_VERSION,
                            {'max_text': MAX_TEXT, 'page_limit': MAX_PAGES, 'max_ocr_pages': MAX_OCR_PAGES, 'adaptive_routing': adaptive.VERSION}) if checkpoint_dir or destination else None
    reused, retriable, attempted, failed, stop_reasons = [], [], [], [], []
    deferred = set(defer_pages)
    # Read all completed checkpoints before spending this batch's time budget.
    # A failed early page must not hide completed later pages on a short retry.
    completed_cache = {n: value for n in range(1, min(count, MAX_PAGES) + 1)
                       if cache and (value := cache.load(n))}
    pending_cache = {n: value for n in range(1, min(count, MAX_PAGES) + 1)
                     if cache and n not in completed_cache and (value := cache.load(n, incomplete=True))}
    characters = sum(len(v['page']['text']) for v in [*completed_cache.values(), *pending_cache.values()])
    ocr_count = 0
    regions, multi_column = [], []
    truncated = False
    text_budget_hit = characters >= MAX_TEXT
    with tempfile.TemporaryDirectory(prefix="zhiheng-pdf-") as folder:
        for index in range(min(count, MAX_PAGES)):
            if cancelled():
                raise InterruptedError('已取消文档解析，已完成页保持可恢复')
            cached = completed_cache.get(index + 1)
            if cached:
                result['pages'].append(cached['page'])
                regions.extend(cached['regions'])
                if cached.get('multi_column'): multi_column.append(index + 1)
                reused.append(index + 1)
                progress('恢复已处理页面',20 + int(65 * (index + 1) / min(count, MAX_PAGES)),f'已恢复第 {index + 1} 页，无需重复解析')
                continue
            remaining = MAX_DOCUMENT_SECONDS - (time.monotonic() - started)
            reason = ('retry_exhausted' if index + 1 in deferred else
                      'time_budget' if remaining <= 0 else
                      'ocr_batch_budget' if MAX_OCR_PAGES > 0 and ocr_count >= MAX_OCR_PAGES else
                      'text_budget' if text_budget_hit or characters >= MAX_TEXT else None)
            if reason:
                truncated = True
                stop_reasons.append(reason)
                previous = pending_cache.get(index + 1)
                if previous:
                    result['pages'].append(previous['page'])
                    regions.extend(previous['regions'])
                    retriable.append(index + 1)
                continue
            number = index + 1
            previous = pending_cache.get(number)
            if previous:
                characters -= len(previous['page']['text'])
            attempted.append(number)
            page_started = time.monotonic()
            ocr_attempted = False
            retry_page = False
            routing = None
            comparison_result = None
            image_path = Path(folder) / "page.png"
            current = page(number, f"第 {number} 页", "")
            try:
                extracted = _pdf_worker(path, index, image_path, min(PDF_PAGE_TIMEOUT, remaining), cancelled=cancelled)
                text, blocks = text_blocks(extracted["blocks"])
                native_tables = extracted["tables"]
                with Image.open(image_path) as image:
                    quality = quality_metrics(image)
                    quality.update(rotation_correction=0, ocr_mean_confidence=None, pdf_page_rotation=extracted["page_rotation"])
                    extraction = "native" if text else "unreadable"
                    page_tables = native_tables
                    native_text, native_blocks = text, blocks
                    routing = adaptive.plan(extracted, native_text, quality)
                    # Exact rendered input used for quality/OCR; this is not a
                    # model-proposed target box or a claim of semantic precision.
                    from hashlib import sha256
                    routing['original_page_image_sha256'] = sha256(image_path.read_bytes()).hexdigest()
                    if routing['requested_route'] in ('ocr_primary', 'native_with_ocr_comparison'):
                        if ocr_count >= MAX_OCR_PAGES:
                            quality['warnings'].append(f'已达到 {MAX_OCR_PAGES} 页 OCR 上限，此页需对照的范围尚未识别；可拆分后重试')
                            routing.update(status='skipped_ocr_page_limit', executed=False)
                            truncated = True
                            retry_page = True
                        else:
                            remaining = MAX_DOCUMENT_SECONDS - (time.monotonic() - started)
                            ocr_started = time.monotonic()
                            try:
                                if remaining <= 0:
                                    raise OCRError('文档解析已达到耗时上限，此页 OCR 未完成')
                                ocr_count += 1
                                ocr_attempted = True
                                routing['ocr_attempted'] = True
                                recognized = recognize(image, timeout=min(35, remaining), cancelled=cancelled)
                                if cancelled():
                                    raise InterruptedError('已取消文档解析，未发布晚到OCR结果')
                                comparison_result = adaptive.comparison(native_text, native_blocks, recognized, routing)
                                # Only absent or geometrically marginal native text may yield
                                # primary OCR. Reliable body text and its offsets stay intact.
                                if routing['requested_route'] == 'ocr_primary' and recognized['text']:
                                    text, blocks, quality, extraction = recognized['text'], recognized['blocks'], recognized['quality'], recognized['extraction']
                                    page_tables = recognized['tables'] or native_tables
                                else:
                                    quality['ocr_comparison_quality'] = recognized['quality']
                                    if recognized['quality'].get('preprocessing'):
                                        quality['preprocessing'] = recognized['quality']['preprocessing']
                                    quality['warnings'].extend(recognized['quality'].get('warnings', []))
                                if native_text and comparison_result['requires_original_page_review']:
                                    quality['visual_review_required'] = True
                                    quality['warnings'].append('原生文字与本地OCR对照存在差异；两者已分别保留，数字/符号及新增图内文字须对照原页核验，不能据置信度覆盖原值')
                                if routing['requested_route'] == 'native_with_ocr_comparison':
                                    quality['warnings'].append('补充OCR为待复核候选，未并入原生正文；原页图表读取可使用独立观测并继续核验')
                                routing.update(status='executed', executed=True,
                                               selected_extraction=extraction,
                                               ocr_result_available=bool(recognized['text']))
                                if routing['requested_route'] == 'ocr_primary' and not recognized['text']:
                                    routing['status'] = 'ocr_primary_empty'
                                    quality['warnings'].append('本页主要内容需要OCR，但本次未识别到正文；保留原生页码/标题并标待重试，不作为完整页面缓存')
                                    retry_page = True
                                    failed.append(number)
                                elif not recognized['text']:
                                    routing['status'] = 'comparison_empty_native_retained'
                            except OCRError as exc:
                                quality['warnings'].append(str(exc))
                                routing.update(status='ocr_failed', executed=False)
                                retry_page = True
                                failed.append(number)
                            finally:
                                routing['ocr_elapsed_ms'] = round((time.monotonic() - ocr_started) * 1000)
                    else:
                        routing.update(status='executed', executed=True, selected_extraction=extraction)
                    quality["warnings"].extend(extracted["warnings"])
                    if not text and not quality["warnings"]:
                        quality["warnings"].append("未提取到可用文字；请查看原件或重新扫描")
                if characters + len(text) > MAX_TEXT:
                    result["warnings"].append(f"已达到 {MAX_TEXT:,} 字符上限，其余页面尚未处理")
                    truncated = True
                    text_budget_hit = True
                    stop_reasons.append('text_budget')
                    # The page is not partially truncated into false evidence.
                    current.update(extraction='unreadable', quality={'warnings': ['本页文字超出文档字符预算，未发布截断正文'], 'blank': False}, blocks=[])
                    if cache:
                        cache.save(current, [], incomplete=True)
                    retriable.append(number)
                    deferred.add(number)
                    continue
                current.update(text=text, char_count=len(text), has_image=extracted["has_image"], blocks=blocks,
                               extraction=extraction, quality=quality,
                               printed_page_label=extracted.get('printed_page_label'),
                               printed_page_label_basis='pdf_page_label' if extracted.get('printed_page_label') else 'unknown')
                if comparison_result is not None:
                    current['ocr_comparison'] = comparison_result
                    current['chart_observations'] = adaptive.chart_observations(comparison_result)
                characters += len(text)
                # Layout diagnostic only; does not manufacture a reading-order guarantee.
                left = [b for b in blocks if b["bbox"][2] < .53 and len(b["text"]) > 15]
                right = [b for b in blocks if b["bbox"][0] > .47 and len(b["text"]) > 15]
                if len(left) >= 3 and len(right) >= 3:
                    multi_column.append(number)
                    quality["warnings"].append("页面可能包含分栏，文字按区域位置整理，跨栏阅读顺序需对照原件")
                for table in page_tables:
                    table.update(table_name=f"第{number}页_表{sum(r['page'] == number for r in regions) + 1}", page=number)
                    regions.append(table)
            except (ParseError, OCRError) as exc:
                if previous and previous['page']['text']:
                    # Retain an earlier exact native/OCR observation if a later
                    # retry fails before reading this page. It remains pending.
                    import copy
                    current = copy.deepcopy(previous['page'])
                    current.setdefault('quality', {}).setdefault('warnings', []).append('本次重试失败：' + str(exc))
                    current['retry_status'] = 'failed_previous_observation_retained'
                    regions.extend(copy.deepcopy(previous['regions']))
                    characters += len(current['text'])
                else:
                    current.update(extraction="unreadable", blocks=[], quality={"warnings": [str(exc)], "blank": False, "rotation_correction": 0, "ocr_mean_confidence": None})
                retry_page = True
                failed.append(number)
            if cancelled():
                raise InterruptedError('已取消文档解析，未发布晚到页面')
            if routing is not None:
                routing['elapsed_ms'] = round((time.monotonic() - page_started) * 1000)
                current['adaptive_processing'] = routing
            current['processing_history'] = [{'method': 'pdf_page_pipeline', 'parser_version': PARSER_VERSION,
                'elapsed_ms': round((time.monotonic() - page_started) * 1000), 'ocr_attempted': ocr_attempted,
                'selected_extraction': current.get('extraction', 'unreadable'),
                'status': 'unreadable' if current.get('extraction') == 'unreadable' else 'extracted',
                'verification_status': 'not_verified'}]
            result["pages"].append(current)
            if current.get('extraction') == 'unreadable' and not current.get('quality', {}).get('blank'):
                retry_page = True
                failed.append(number)
            if retry_page:
                retriable.append(number)
                if cache:
                    cache.save(current, [r for r in regions if r['page'] == number], number in multi_column, incomplete=True)
            elif cache:
                cache.save(current, [r for r in regions if r['page'] == number], number in multi_column)
            progress("读取 PDF 页面", 20 + int(65 * number / min(count, MAX_PAGES)), f"已处理 {number}/{count} 页")
    if count > MAX_PAGES:
        truncated = True
        result["warnings"].append(f"原件共 {count} 页，本次处理上限 {MAX_PAGES} 页")
        stop_reasons.append('page_policy')
    if 'time_budget' in stop_reasons:
        result['warnings'].append(f'本批达到 {MAX_DOCUMENT_SECONDS} 秒预算，已完成页可复用，剩余页尚未处理，需继续解析')
    if 'ocr_batch_budget' in stop_reasons:
        result['warnings'].append(f'本批达到 {MAX_OCR_PAGES} 页 OCR 预算，已完成页可复用，剩余页需继续处理')
    for current in result["pages"]:
        result["warnings"].extend(f"第 {current['number']} 页：{notice}" for notice in current.get("quality", {}).get("warnings", []))
    characters = sum(len(p['text']) for p in result['pages'])
    result.update(page_count=count, summary=f"{count} 页 · 已提取 {characters:,} 字符")
    result["metadata"] = dict(page_unit="physical_page", parsed_pages=len(result["pages"]), text_characters=characters,
                              ocr_performed=bool(ocr_count or any(p.get('extraction') == 'ocr' or p.get('ocr_comparison') for p in result['pages'])), outline=info["outline"], outline_note="使用 PDF 内嵌书签" if info["outline"] else "原件没有内嵌书签；未从正文推断完整章节目录", quality_summary=_quality_metadata(result["pages"], truncated or bool(retriable)),
                              complexity={"multi_column_pages": multi_column, "merged_table_count": sum(r["has_merged_cells"] for r in regions),
                                          "formula_candidates": sum(b["kind"] == "formula_candidate" for p in result["pages"] for b in p.get("blocks", [])),
                                          "notes": ["公式区域保留原文字和坐标，未验证数学含义。", "仅将闭合格线支持的表格转为数据；无框、破损或倾斜扫描表保留文字供核对。"]},
                              extraction_limits={"max_pages": MAX_PAGES, "max_ocr_pages": MAX_OCR_PAGES, "max_document_seconds": MAX_DOCUMENT_SECONDS,
                                                 "max_page_seconds": PDF_PAGE_TIMEOUT, "max_image_pixels": 40_000_000})
    pending = sorted(set(range(1,min(count,MAX_PAGES)+1)) - {p['number'] for p in result['pages']} | set(retriable))
    result['metadata']['import_batches'] = {'schema_version':1,'source_sha256':cache.source_sha256 if cache else None,
        'parser_version':PARSER_VERSION,'reused_pages':reused,'pending_pages':pending,
        'completed_pages':[p['number'] for p in result['pages'] if p['number'] not in retriable],
        'resumable':bool(cache and pending and 'text_budget' not in stop_reasons),'status':'partial' if pending or count > MAX_PAGES else 'complete',
        'out_of_policy_pages':max(0,count-MAX_PAGES),'model_requests_replayed':False,
        'attempted_pages':attempted, 'failed_pages':sorted(set(failed)),
        'stop_reasons':sorted(set(stop_reasons)), 'ocr_attempts':ocr_count,
        'completeness_note':'complete表示策略范围内页面解析结束，不表示OCR、图表或问答正确率已核验'}
    if destination is not None and not defer_table_write:
        _save_regions(result, regions, destination)
    else:
        result["metadata"]["table_regions"] = (regions if defer_table_write else
            [{k: v for k, v in r.items() if k != "rows"} for r in regions])
        _region_semantics(result, regions)
    return document_structure.enrich(result, format='pdf')


def parse_pdf_batches(path, progress, destination=None, checkpoint_dir=None, cancelled=lambda: False, *, max_batches=16):
    """Continue local parsing with bounded retries; never replay a model request.

    Each batch keeps the existing time/OCR safeguards. Successful pages are
    checkpointed once, persistently failed pages remain visible, and a stalled
    or policy-limited document is returned as partial instead of looping.
    """
    if type(max_batches) is not int or not 1 <= max_batches <= 32:
        raise ValueError('PDF 自动续批次数须为1到32')
    if checkpoint_dir is None and destination is None:
        with tempfile.TemporaryDirectory(prefix='zhiheng-page-checkpoints-') as temporary:
            result = parse_pdf_batches(path, progress, checkpoint_dir=temporary, cancelled=cancelled, max_batches=max_batches)
            result['metadata']['import_batches'].update(resumable=False, checkpoint_scope='temporary_call_only')
            return result
    failure_counts, summaries, previous_completed = {}, [], set()
    result = None
    stop = 'batch_limit'
    for batch in range(1, max_batches + 1):
        if cancelled():
            raise InterruptedError('已取消文档续批，已完成页保持可恢复')
        exhausted = {n for n, attempts in failure_counts.items() if attempts >= 2}
        progress('继续解析文档', 20, f'第 {batch} 批：复用完成页，处理尚未完成的页面')
        result = parse_pdf(path, progress, destination, checkpoint_dir, cancelled,
                           defer_pages=exhausted, defer_table_write=True)
        state = result['metadata']['import_batches']
        for number in state['failed_pages']:
            failure_counts[number] = failure_counts.get(number, 0) + 1
        completed = set(state['completed_pages'])
        summaries.append({'batch': batch, 'new_completed_pages': sorted(completed - set(state['reused_pages'])),
                          'attempted_pages': state['attempted_pages'], 'failed_pages': state['failed_pages'],
                          'pending_count': len(state['pending_pages']), 'stop_reasons': state['stop_reasons'],
                          'ocr_attempts': state['ocr_attempts']})
        if state['status'] == 'complete':
            stop = 'complete'
            break
        if 'text_budget' in state['stop_reasons']:
            stop = 'text_budget'
            break
        eligible = set(state['pending_pages']) - {n for n, attempts in failure_counts.items() if attempts >= 2}
        if not eligible:
            stop = 'retry_exhausted' if state['pending_pages'] else 'page_policy'
            break
        if not state['attempted_pages'] or (completed == previous_completed and not state['failed_pages']):
            stop = 'no_progress'
            break
        previous_completed = completed
    state = result['metadata']['import_batches']
    state.update(automatic_batches=summaries, batches_run=len(summaries), max_batches=max_batches,
                 batch_stop_reason=stop, failed_attempts={str(k): v for k, v in sorted(failure_counts.items())},
                 exhausted_pages=sorted(n for n, attempts in failure_counts.items() if attempts >= 2))
    if state['status'] != 'complete':
        result['warnings'].append(f"自动续批已停止（{stop}）；仍有 {len(state['pending_pages'])} 页待处理、{state['out_of_policy_pages']} 页超出策略范围，请查看页面状态后重试或拆分原件")
    # Materialize tables once after all batches. Repeated CREATE TABLE on the
    # same scratch database would otherwise fail at the second batch.
    regions = result['metadata']['table_regions']
    if destination is not None:
        _save_regions(result, regions, destination)
    else:
        result['metadata']['table_regions'] = [{k: v for k, v in region.items() if k != 'rows'} for region in regions]
    return result


def parse_image(path, destination, progress, cancelled=lambda: False):
    from PIL import Image, ImageOps
    from .ocr_native import OCRError, normalized_image
    from .ocr_adapter import recognize
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as original:
            if original.width * original.height > 40_000_000:
                raise ParseError("图片超过 4,000 万像素，请缩小后导入")
            original.verify()
        with Image.open(path) as original:
            orientation = original.getexif().get(274, 1)
            image = normalized_image(original)
            width, height = image.size
            progress("本地文字识别", 40, "正在识别图片文字与位置")
            try:
                recognized = recognize(image, cancelled=cancelled)
            except OCRError as exc:
                raise ParseError(str(exc)) from exc
    result = base_result()
    current = page(1, "原始图片", recognized["text"], True)
    current.update(blocks=recognized["blocks"], quality=recognized["quality"], extraction=recognized["extraction"])
    result.update(pages=[current], page_count=1, summary=f"{width} × {height} 像素 · {len(recognized['text']):,} 字符", warnings=recognized["quality"]["warnings"][:])
    result["metadata"] = dict(width=width, height=height, exif_orientation=orientation, ocr_performed=not recognized["quality"]["blank"],
                              page_unit="physical_page", outline=[], quality_summary=_quality_metadata([current]),
                              complexity={"multi_column_pages": [], "merged_table_count": 0, "formula_candidates": sum(b["kind"] == "formula_candidate" for b in current["blocks"]),
                                          "notes": ["OCR 仅保留公式候选文字，未验证公式。", "无可靠闭合格线时不生成数据表。"]})
    for i, region in enumerate(recognized["tables"]):
        region.update(table_name=f"图片表{i + 1}", page=1)
    _save_regions(result, recognized["tables"], destination)
    return document_structure.enrich(result, format=path.suffix.lower().lstrip('.'))



def word_table_grid(node, ns):
    """Expand Word grid positions without inventing values for merged coverage."""
    attribute = "{" + ns["w"] + "}val"
    declared_width = len(node.findall("w:tblGrid/w:gridCol", ns))
    rows, merges, active, notes, cells, header_rows = [], [], {}, [], [], []
    reliable = True
    plain_rows = []
    plain_offset = 0
    for row_index, row in enumerate(node.findall("w:tr", ns), start=1):
        if row_index > MAX_ROWS or row_index * max(1, declared_width) > MAX_CELLS:
            raise ParseError('Word 表格超出行数或单元格预算，请拆分后导入')
        header_flag = row.find('w:trPr/w:tblHeader', ns)
        if header_flag is not None and header_flag.get(attribute, 'true').lower() not in ('0', 'false', 'off'):
            header_rows.append(row_index)
        before = row.find("w:trPr/w:gridBefore", ns)
        after = row.find("w:trPr/w:gridAfter", ns)
        left = int(before.get(attribute, "0")) if before is not None else 0
        right = int(after.get(attribute, "0")) if after is not None else 0
        if not 0 <= left <= MAX_COLUMNS or not 0 <= right <= MAX_COLUMNS:
            raise ParseError("Word 表格的空白网格超出处理范围")
        values, used, plain = [None] * left, set(), []
        for cell in row.findall("w:tc", ns):
            text = "".join(t.text or "" for t in cell.findall(".//w:t", ns))
            text_start = plain_offset + sum(len(value) + 3 for value in plain)
            plain.append(text)
            span_node = cell.find("w:tcPr/w:gridSpan", ns)
            span = int(span_node.get(attribute, "1")) if span_node is not None else 1
            if span < 1 or len(values) + span > MAX_COLUMNS:
                raise ParseError("Word 表格合并跨度超出处理范围")
            start = len(values)
            if cell.find('w:tcPr/w:hMerge', ns) is not None:
                reliable = False
                notes.append('存在旧式横向合并标记，未补猜其网格覆盖范围')
            vertical = cell.find("w:tcPr/w:vMerge", ns)
            mode = vertical.get(attribute, "continue") if vertical is not None else None
            origin = None
            display_text = text
            if mode == "restart":
                origin = {"row": row_index, "column": start + 1}
                for column in range(start, start + span):
                    active[column] = origin
                    used.add(column)
            elif mode == "continue":
                prior = [active.get(column) for column in range(start, start + span)]
                if not prior or prior[0] is None or any(item != prior[0] for item in prior):
                    reliable = False
                    notes.append("纵向合并单元格缺少可确认的起点")
                else:
                    origin = prior[0]
                    original_columns = {column for column, item in active.items() if item == origin}
                    if original_columns != set(range(start, start + span)):
                        reliable = False
                        notes.append('纵向合并续接与起点的列跨度不一致')
                used.update(range(start, start + span))
                display_text = None
            elif mode is not None:
                reliable = False
                notes.append("存在不支持的纵向合并标记")
            if span > 1 or vertical is not None:
                merges.append({"row": row_index, "column": start + 1, "column_span": span, "vertical": mode, "origin": origin})
            covered = [origin['row'], origin['column']] if mode == 'continue' and origin else None
            cells.append({'row': row_index, 'col': start + 1, 'text': display_text,
                          'rowspan': 1, 'colspan': span, 'covered_by': covered,
                          'start': text_start, 'end': text_start + len(text),
                          **({'source_text': text} if mode == 'continue' and text else {})})
            for column in range(start + 2, start + span + 1):
                cells.append({'row': row_index, 'col': column, 'text': None, 'rowspan': 1, 'colspan': 1,
                              'covered_by': covered or [row_index, start + 1]})
            values.extend([display_text] + [None] * (span - 1))
        active = {column: origin for column, origin in active.items() if column in used}
        values.extend([None] * right)
        if declared_width and len(values) > declared_width:
            reliable = False
            notes.append("实际单元格跨度超过声明的表格网格")
        rows.append(values)
        if len(cells) > MAX_CELLS:
            raise ParseError('Word 表格超出单元格预算，请拆分后导入')
        plain_rows.append(" | ".join(plain))
        plain_offset += len(plain_rows[-1]) + 1
    width = declared_width or max((len(row) for row in rows), default=0)
    if width > MAX_COLUMNS:
        raise ParseError(f"Word 表格超过 {MAX_COLUMNS} 列")
    if node.findall(".//w:tc/w:tbl", ns):
        notes.append("含嵌套表格，内层内容以单元格文字展示")
        reliable = False
    anchors = {(c['row'], c['col']): c for c in cells}
    for cell in cells:
        if cell.get('covered_by'):
            anchor = anchors.get(tuple(cell['covered_by']))
            if anchor:
                anchor['rowspan'] = max(anchor['rowspan'], cell['row'] - anchor['row'] + 1)
    return {"rows": rows, "width": width, "merges": merges, "reliable": reliable,
            "cells": cells, "header_rows": header_rows, "notes": list(dict.fromkeys(notes)), "plain": "\n".join(plain_rows)}

def parse_docx(path, destination, progress):
    check_zip(path)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        try:
            raw = archive.read("word/document.xml")
        except KeyError as exc:
            raise ParseError("Word 文档缺少正文，请重新导出 DOCX") from exc
        if len(raw) > 32 * 1024 * 1024:
            raise ParseError("Word 正文结构过大，请拆分后导入")
        if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
            raise ParseError("Word 文档含不受支持的实体声明")
        root = ET.fromstring(raw)
        try:
            styles = document_structure.docx_styles(archive.read('word/styles.xml'))
        except KeyError:
            styles = {}
        image_count = sum(i.filename.startswith("word/media/") for i in archive.infolist())
        body = root.find("w:body", ns)
        marginal_parts = document_structure.docx_marginal_parts(archive, body) if body is not None else {'parts': [], 'unresolved_references': []}
    body = root.find("w:body", ns)
    if body is None:
        raise ParseError("Word 文档没有可读取的正文")
    result, chunks, title, text_count = base_result(), [], "正文", 0
    writer = TableWriter(destination)
    table_count = 0
    layouts = []
    semantic_specs, preceding, pending_spec = [], [], None
    outline = []
    formula_count = 0
    formulas, paragraphs, formula_bytes = [], [], 0
    content_controls = []
    section_chars = 0
    def save_formula(math_node, location):
        nonlocal formula_bytes
        if len(formulas) >= formula_structure.MAX_FORMULAS or formula_bytes >= 2_000_000:
            return
        value = formula_structure.from_omml(math_node, identifier=f'F{len(formulas) + 1}',
            location={**location, 'archive_member': 'word/document.xml'})
        size = len(json.dumps(value, ensure_ascii=False))
        if formula_bytes + size > 2_000_000:
            formula_bytes = 2_000_000
            return
        formulas.append(value)
        formula_bytes += size
    try:
        for body_index, node, block_path, container_toc, controls in document_structure.docx_body_blocks(body, with_controls=True):
            if node.tag.endswith("}p"):
                text = "".join(t.text or "" for t in node.findall(".//w:t", ns)).strip()
                math_ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
                math_nodes = node.findall('.//m:oMath', math_ns)
                math_tokens = [" ".join(t.text or "" for t in m.findall(".//m:t", math_ns)).strip()
                               for m in math_nodes]
                if math_tokens:
                    formula_count += len(math_tokens)
                    # Tokens are preserved, not silently flattened into a verified equation.
                    text += "\n" + "\n".join("[公式结构待核对] " + value for value in math_tokens if value)
                properties = document_structure.paragraph_properties(node, styles)
                if container_toc:
                    properties.update(is_toc=True, heading_level=None)
                if text and properties['heading_level']:
                    pending_spec = None
                    preceding = []
                    if chunks:
                        result["pages"].append(page(len(result["pages"]) + 1, title, "\n".join(chunks)))
                        chunks = []
                        section_chars = 0
                    title = text
                    outline.append({"level": properties['heading_level'],
                                    "title": text, "page": len(result["pages"]) + 1, "source": "docx_heading"})
                if text:
                    location = {'page': len(result['pages']) + 1, 'start': section_chars,
                                'page_unit': 'logical_section', 'body_index': body_index, 'block_path': block_path}
                    location['end'] = location['start'] + len(text)
                    content_controls.extend(document_structure.docx_control_annotations(node, text, location, controls))
                    if len(paragraphs) < document_structure.MAX_RECORDS:
                        paragraphs.append({'text': text, 'source_position': dict(location), **properties})
                    token_cursor = 0
                    for math_index, (math_node, tokens) in enumerate(zip(math_nodes, math_tokens)):
                        marker = '[公式结构待核对] ' + tokens
                        found = text.find(marker, token_cursor) if tokens else -1
                        position = {**location, 'math_index': math_index, 'text_span_contains_formula': found >= 0}
                        if found >= 0:
                            position.update(start=location['start'] + found, end=location['start'] + found + len(marker))
                            token_cursor = found + len(marker)
                        save_formula(math_node, position)
                    paragraph = {'text': text, 'evidence': location}
                    preceding = (preceding + [paragraph])[-3:]
                    if table_semantics.CAPTION.match(text):
                        pending_spec = None
                        # A new caption starts its own association window. Text
                        # after the previous table cannot become this one's unit.
                        preceding = [paragraph]
                    if pending_spec is not None and len(pending_spec['after']) < 3:
                        pending_spec['after'].append(paragraph)
                    chunks.append(text)
                    section_chars += len(text) + 1
                    text_count += len(text)
            elif node.tag.endswith("}tbl"):
                grid = word_table_grid(node, ns)
                if grid["rows"]:
                    table_count += 1
                    table_name = f"表格 {table_count}"
                    if grid["reliable"]:
                        writer.add(table_name, [f"列{i + 1}" for i in range(grid["width"])], grid["rows"])
                    else:
                        result["warnings"].append(f"{table_name}的网格或合并关系无法可靠确定，已仅保留原文文字，未生成数据表")
                    if grid["merges"]:
                        result["warnings"].append(f"{table_name}含合并单元格；已定位的覆盖位置以空值显示，请结合原件阅读，勿把空值直接视为缺失业务数据")
                    result["warnings"].extend(f"{table_name}：{note}" for note in grid["notes"])
                    layouts.append({"name": table_name, "columns": grid["width"], "reliable_grid": grid["reliable"], "merged_cells": grid["merges"]})
                    start = section_chars
                    position = {'page': len(result['pages']) + 1, 'start': start, 'end': start + len(grid['plain']),
                                'body_index': body_index, 'block_path': block_path, 'page_unit': 'logical_section'}
                    content_controls.extend(document_structure.docx_control_annotations(node, grid['plain'], position, controls))
                    for math_index, math_node in enumerate(node.iter(f'{{{formula_structure.M}}}oMath')):
                        formula_count += 1
                        save_formula(math_node, {**position, 'table_name': table_name, 'math_index': math_index,
                                                 'text_span_contains_formula': False})
                    if len(semantic_specs) < table_semantics.MAX_TABLES:
                        semantic_cells = []
                        for cell in grid['cells']:
                            semantic_cells.append({**cell, **{k: cell[k] + start for k in ('start', 'end') if k in cell}})
                        pending_spec = {'region': {'table_name': table_name, 'row_count': len(grid['rows']),
                            'column_count': grid['width'], 'cells': semantic_cells, 'requires_review': not grid['reliable'],
                            'review_reasons': grid['notes']}, 'source_position': position,
                            'explicit_header_rows': grid['header_rows'], 'before': preceding, 'after': []}
                        semantic_specs.append(pending_spec)
                    else:
                        pending_spec = None
                    preceding = []
                    chunks.append(grid["plain"])
                    section_chars += len(grid['plain']) + 1
                    text_count += len(grid["plain"])
            if len(content_controls) > document_structure.MAX_RECORDS:
                raise ParseError('Word 内容控件超过10000条，无法完整保留标签，请拆分后导入')
            if text_count > MAX_TEXT or len(result["pages"]) >= MAX_PAGES and chunks:
                raise ParseError("Word 正文超出本阶段处理范围，请按章节拆分后导入")
        if chunks:
            result["pages"].append(page(len(result["pages"]) + 1, title, "\n".join(chunks), bool(image_count)))
    finally:
        writer.close()
    result["tables"], _ = sqlite_tables(destination, observe_types=True)
    result["data_path"] = str(destination) if table_count else None
    result["page_count"] = len(result["pages"])
    result["summary"] = f"{len(result['pages'])} 个逻辑节 · {len(result['tables'])} 张表"
    result["metadata"] = dict(page_unit="logical_section", pagination_note="DOCX 按标题划分逻辑节，不代表 Word 打印页码", embedded_images=image_count, text_characters=text_count, table_layouts=layouts)
    result['metadata'].update(paragraphs=paragraphs, paragraphs_truncated=len(paragraphs) >= document_structure.MAX_RECORDS,
                              formula_structures=formulas, formulas_seen=formula_count,
                              formula_structures_truncated=formula_count > len(formulas),
                              content_controls=content_controls, docx_marginal_parts=marginal_parts,
                              content_control_note='tag和alias是原件声明，不自动视为已核验业务事实；只有span_verified条目具备正文精确字符锚点')
    result['metadata']['table_semantics'] = [table_semantics.build_table_semantics(format='docx', **spec) for spec in semantic_specs]
    result['metadata']['table_semantics_truncated'] = table_count > len(semantic_specs)
    if table_count > len(semantic_specs):
        result['warnings'].append(f'仅为前{table_semantics.MAX_TABLES}张表保存语义关系，其余表保留原始文字和网格')
    result["metadata"].update(outline=outline, complexity={"multi_column_pages": [], "merged_table_count": sum(bool(t["merged_cells"]) for t in layouts),
        "formula_candidates": formula_count, "notes": ["DOCX 公式保留原始词元及有界OMML结构；结构解析不等于语义核验，未核验表达式不得计算。"]})
    for section in result["pages"]:
        section["extraction"] = "native"
        section["blocks"] = []
        offset = 0
        for line in section["text"].split("\n"):
            section["blocks"].append(dict(text=line, bbox=None, confidence=None, start=offset, end=offset + len(line),
                                           kind="formula_candidate" if line.startswith("[公式结构待核对]") else "text"))
            offset += len(line) + 1
        section["quality"] = {"warnings": [], "coordinate_note": "DOCX 逻辑节不提供物理页坐标"}
    result["metadata"]["quality_summary"] = _quality_metadata(result["pages"])
    if formula_count:
        result["warnings"].append(f"发现 {formula_count} 处 Word 公式，已保留原始词元和 {len(formulas)} 处OMML结构；公式结构仍需对照原件")
    if formula_count > len(formulas):
        result['warnings'].append('公式结构元数据达到数量或字符预算，其余原文仍保留，未解析结构明确为未知')
    if image_count:
        result["warnings"].append(f"文档含 {image_count} 个嵌入图片，尚未进行图片文字识别")
    if not text_count:
        result["warnings"].append("文档未提取到正文文字，待文字识别；可以下载原件")
    progress("整理 Word 结构", 85, f"已整理 {len(result['pages'])} 个逻辑节")
    return document_structure.enrich(result, format='docx')


def _xlsx_semantics(path, sheet, sampled, header_row, row_runs, imported_rows, separated):
    """Read merge declarations without loading an unbounded editable workbook."""
    from openpyxl.utils.cell import range_boundaries
    merges, merge_truncated = [], False
    member = sheet._worksheet_path.lstrip('/')
    with zipfile.ZipFile(path) as archive, archive.open(member) as stream:
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag.endswith('}mergeCell'):
                if len(merges) >= 4096:
                    merge_truncated = True
                else:
                    try:
                        c1, r1, c2, r2 = range_boundaries(element.get('ref', ''))
                        if not (1 <= r1 <= r2 <= MAX_ROWS + 1 and 1 <= c1 <= c2 <= MAX_COLUMNS):
                            raise ValueError()
                        merges.append((r1, c1, r2, c2))
                    except (ValueError, TypeError):
                        merge_truncated = True
            element.clear()
    cells, after = [], []
    for row, values in sorted(sampled.items()):
        for col, value in enumerate(values, 1):
            cell = {'row': row, 'col': col, 'text': value.isoformat() if hasattr(value, 'isoformat') else value,
                    'rowspan': 1, 'colspan': 1, 'covered_by': None}
            matching = [m for m in merges if m[0] <= row <= m[2] and m[1] <= col <= m[3]]
            if len(matching) > 1:
                merge_truncated = True
            if matching:
                r1, c1, r2, c2 = matching[0]
                if (row, col) == (r1, c1):
                    cell.update(rowspan=r2-r1+1, colspan=c2-c1+1)
                else:
                    cell.update(text=None, covered_by=[r1, c1])
            cells.append(cell)
            if isinstance(value, str) and table_semantics.FOOTNOTE.match(value):
                after.append({'text': value, 'evidence': {'sheet': sheet.title, 'row': row, 'col': col}})
    reasons = []
    if separated:
        reasons.append('工作表含空行分隔的多个数据段，不能自动视为共享表头或单位的单表')
    if merge_truncated:
        reasons.append('工作表合并声明不完整或重叠，合并结构需复核')
    region = {'table_name': sheet.title, 'row_count': sheet.max_row or max(sampled, default=0),
              'column_count': sheet.max_column or max((len(v) for v in sampled.values()), default=0),
              'cells': cells, 'requires_review': True, 'review_reasons': reasons,
              'truncated': merge_truncated or (sheet.max_row or 0) > len(sampled)}
    result = table_semantics.build_table_semantics(region, format='xlsx', source_position={'sheet': sheet.title}, after=after)
    result['merged_ranges'] = [{'start_row': r1, 'start_col': c1, 'end_row': r2, 'end_col': c2} for r1, c1, r2, c2 in merges]
    result['merge_ranges_truncated'] = merge_truncated
    result['sql_row_mapping'] = {'header_source_row': header_row, 'sql_row_count': imported_rows,
        'runs': row_runs, 'run_format': '[first_sql_row,first_source_row,count]',
        'note': 'SQL保留首个非空行作列名及之后非空行；语义候选不改写已导入值，源坐标不等同SQL行号。'}
    if separated:
        # The SQL import intentionally preserves its historic single-sheet
        # representation. A blank-separated segment cannot inherit the first
        # segment's semantic header, even when its SQL column name is unchanged.
        result['structure_unknown'] = True
        result['header_basis'] = 'unknown'
        result['header_rows'] = []
        result['row_headers'] = []
        for column in result['columns']:
            column.update(header_path=[], label='', role_basis='unknown')
        for unit in result['units']:
            unit['scope'] = 'ambiguous'
    return result


def parse_table(path, fmt, destination, progress):
    result = base_result()
    writer = TableWriter(destination)
    try:
        if fmt == "csv":
            text, encoding = decode_text(path.read_bytes())
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
            headers = next(rows, [])
            if not headers:
                raise ParseError("CSV 文件为空，至少需要一行表头")
            writer.add("数据表", headers, (row for row in rows if any(x.strip() for x in row)))
            result["metadata"] = dict(encoding=encoding, header_note="首行作为列名；空白或重复列名已规范化", cell_types="CSV 保留原始字符串，未推断数值类型")
        else:
            from openpyxl import load_workbook
            check_zip(path)
            # Both cached values and formulas are inspected; formulas are never evaluated.
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            formula_count, missing_cache = 0, 0
            semantics = []
            try:
                if len(book.worksheets) > 128:
                    raise ParseError("工作簿超过 128 个工作表，请拆分后导入")
                for index, sheet in enumerate(book.worksheets):
                    if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > MAX_COLUMNS:
                        raise ParseError(f"工作表「{sheet.title}」超出 {MAX_ROWS:,} 行或 {MAX_COLUMNS} 列，请删除空白格式区或拆分后导入")
                    all_rows = enumerate(sheet.iter_rows(values_only=True), 1)
                    first = next(((number, r) for number, r in all_rows if any(v is not None for v in r)), None)
                    if first is None:
                        result["warnings"].append(f"工作表「{sheet.title}」为空，已略过")
                        continue
                    header_row, header = first
                    sampled, tail, runs = {header_row: list(header)}, [], []
                    capture_rows = max(1, min(256, table_semantics.MAX_CELLS // max(1, len(header))))
                    row_state = {'count': 0, 'blank': False, 'separated': False}
                    def data_rows():
                        for number, row in all_rows:
                            if not any(v is not None for v in row):
                                row_state['blank'] = True
                                continue
                            if row_state['blank']:
                                row_state['separated'] = True
                            row_state['blank'] = False
                            row_state['count'] += 1
                            if len(sampled) < capture_rows:
                                sampled[number] = list(row)
                            tail.append((number, list(row)))
                            if len(tail) > 4:
                                tail.pop(0)
                            if runs and number == runs[-1][1] + runs[-1][2]:
                                runs[-1][2] += 1
                            else:
                                runs.append([row_state['count'], number, 1])
                            yield list(row)
                    writer.add(sheet.title, header, data_rows())
                    if len(semantics) < table_semantics.MAX_TABLES:
                        sampled.update(tail)
                        semantics.append(_xlsx_semantics(path, sheet, sampled, header_row, runs, row_state['count'], row_state['separated']))
                    for formula_row, value_row in zip(formulas[sheet.title].iter_rows(), sheet.iter_rows()):
                        for fc, vc in zip(formula_row, value_row):
                            if fc.data_type == "f":
                                formula_count += 1
                                missing_cache += vc.value is None
                    progress("读取工作表", 25 + int(55 * (index + 1) / len(book.worksheets)), f"已读取 {index + 1}/{len(book.worksheets)} 个工作表")
            finally:
                book.close()
                formulas.close()
            result["metadata"] = dict(header_note="每张表首个非空行作为列名；空白或重复列名已规范化。多级表头关系另存语义候选，未改写SQL列和值", formula_cells=formula_count, formulas_evaluated=False,
                                      table_semantics=semantics, table_semantics_truncated=len(book.worksheets) > len(semantics))
            if formula_count:
                result["warnings"].append(f"发现 {formula_count} 个公式单元格，仅显示文件已有缓存值，未重新计算" + (f"；其中 {missing_cache} 个没有缓存值，显示为空" if missing_cache else ""))
    finally:
        writer.close()
    result["tables"], notices = sqlite_tables(destination, observe_types=True)
    result["warnings"].extend(notices)
    if not result["tables"]:
        result["warnings"].append("未发现可读取的数据表")
    result["data_path"] = str(destination)
    result["metadata"]["column_type_note"] = "列类型根据前 200 行非空值观察，不改变原始值，也不代表业务语义"
    result["summary"] = f"{len(result['tables'])} 张表 · {sum(t['row_count'] for t in result['tables']):,} 行"
    return result


def parse(path, fmt, destination, progress, *, checkpoint_dir=None, cancelled=lambda: False, pdf_batch_limit=1):
    started = time.monotonic()
    if fmt not in SUPPORTED:
        raise ParseError("暂不支持此格式。请使用 PDF、DOCX、TXT、MD、CSV、XLSX、PNG、JPG 或 SQLite")
    if path.stat().st_size > MAX_BYTES:
        raise ParseError("文件超过 80 MB，请拆分后导入")
    if path.stat().st_size == 0:
        raise ParseError("文件为空，请检查原文件后重新导入")
    if fmt == "pdf":
        result = (parse_pdf_batches(path, progress, destination, checkpoint_dir, cancelled, max_batches=pdf_batch_limit)
                  if pdf_batch_limit != 1 else parse_pdf(path, progress, destination, checkpoint_dir, cancelled))
    elif fmt == "docx":
        result = parse_docx(path, destination, progress)
    elif fmt in {"csv", "xlsx"}:
        result = parse_table(path, fmt, destination, progress)
    elif fmt in {"sqlite", "db"}:
        result = base_result()
        result["tables"], result["warnings"] = sqlite_tables(path)
        result["data_path"] = str(path)
        result["metadata"] = dict(read_only=True, original_is_copy=True, views_executed=False)
        if not result["tables"]:
            result["warnings"].append("没有可读取的普通数据表")
        result["summary"] = f"{len(result['tables'])} 张只读表 · {sum(t['row_count'] for t in result['tables']):,} 行"
    elif fmt in {"png", "jpg", "jpeg"}:
        result = parse_image(path, destination, progress, cancelled)
    else:
        text, encoding = decode_text(path.read_bytes())
        if len(text) > MAX_TEXT:
            raise ParseError(f"文本超过 {MAX_TEXT:,} 字符，请拆分后导入")
        result = base_result()
        result.update(pages=[page(1, "文本内容", text)], page_count=1, summary=f"{len(text):,} 字符")
        result["metadata"] = dict(encoding=encoding, page_unit="logical_section", text_characters=len(text))
        if not text.strip():
            result["warnings"].append("文本仅包含空白内容，暂无可用于提问的文字")
    progress("保存解析结果", 95, "已读取原件，正在保存可预览内容")
    if 'document_structure' not in result['metadata']:
        document_structure.enrich(result, format=fmt)
    result['metadata']['processing_history'] = [{'method': f'{fmt}_import', 'parser_version': PARSER_VERSION,
        'elapsed_ms': round((time.monotonic() - started) * 1000), 'status': 'parsed',
        'processing_truncated': bool(result['metadata'].get('quality_summary', {}).get('processing_truncated'))}]
    return result


def render_preview(path, fmt, number, output):
    if fmt == "pdf":
        import pymupdf
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ParseError("PDF 已加密，无法预览")
            if number < 1 or number > len(document):
                raise IndexError(number)
            current = document[number - 1]
            scale = min(2, 1600 / max(current.rect.width, current.rect.height))
            current.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(output)
    elif fmt in {"png", "jpg", "jpeg"}:
        if number != 1:
            raise IndexError(number)
        from PIL import Image, ImageOps
        with Image.open(path) as original:
            if original.width * original.height > 40_000_000:
                raise ParseError("图片尺寸超过安全预览范围")
            image = ImageOps.exif_transpose(original)
            image.thumbnail((1600, 1600))
            image.convert("RGBA" if "A" in image.getbands() else "RGB").save(output, format="PNG")
    else:
        raise ParseError("该格式没有图片预览，请查看文字、表格或下载原件")
