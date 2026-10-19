"""Source-coordinate table contracts; synthetic cases never claim OCR accuracy."""
import json
import sqlite3
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook

from backend import parsers, table_semantics as ts


def paragraph(text):
    return f'<w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p>'


def cell(text, props=''):
    return f'<w:tc><w:tcPr>{props}</w:tcPr>{paragraph(text)}</w:tc>'


def row(*cells, header=False):
    return '<w:tr>' + ('<w:trPr><w:tblHeader/></w:trPr>' if header else '') + ''.join(cells) + '</w:tr>'


def table(*rows, width=3):
    return '<w:tbl><w:tblGrid>' + '<w:gridCol/>' * width + '</w:tblGrid>' + ''.join(rows) + '</w:tbl>'


def docx(tmp_path, content):
    path = tmp_path / 'fixture.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + content + '</w:body></w:document>')
    return parsers.parse(path, 'docx', tmp_path / 'parsed.sqlite', lambda *args: None)


def test_docx_multilevel_headers_exact_offsets_units_footnote_and_merged_null(tmp_path):
    content = paragraph('收入') + paragraph('表1 地区收入') + paragraph('单位：万元') + table(
        row(cell('地区', '<w:vMerge w:val="restart"/>'), cell('收入', '<w:gridSpan w:val="2"/>'), header=True),
        row(cell('', '<w:vMerge/>'), cell('本期（万元）'), cell('同期（万元）'), header=True),
        row(cell('同名地区'), cell('0'), cell('5')),
        row(cell('同名地区'), cell('10'), cell('20')),
    ) + paragraph('注：本期为初步统计数。') + paragraph('表2 从业人数') + paragraph('单位：人') + table(
        row(cell('地区'), cell('人数'), header=True), row(cell('同名地区'), cell('30')), width=2)
    result = docx(tmp_path, content)
    first, second = result['metadata']['table_semantics']
    assert first['header_rows'] == [1, 2] and first['header_basis'] == 'explicit'
    assert [p['text'] for p in first['columns'][1]['header_path']] == ['收入', '本期（万元）']
    cells = {c['cell_id']: c for c in first['cells']}
    assert cells['R1C3']['text'] is None and cells['R1C3']['covered_by'] == [1, 2]
    assert cells['R2C1']['text'] is None and cells['R1C1']['rowspan'] == 2
    assert cells['R3C2']['text'] == '0'
    text = result['pages'][0]['text']
    for c in first['cells']:
        if c.get('text') is not None:
            assert text[c['start']:c['end']] == c['text']
    assert cells['R3C1']['start'] != cells['R4C1']['start']
    assert first['caption']['text'] == '表1 地区收入'
    assert first['footnotes'][0]['text'] == '注：本期为初步统计数。'
    assert {u['value'] for u in first['units']} == {'万元'}
    assert {u['value'] for u in second['units']} == {'人'}
    for record in [first['caption'], *first['footnotes'], *first['units']]:
        pos = record['evidence']
        assert text[pos['start']:pos['end']] == record['text']
    with sqlite3.connect(result['data_path']) as database:
        assert database.execute('SELECT * FROM "表格 1" WHERE rowid=3').fetchone() == ('同名地区', '0', '5')


def test_orphan_docx_merge_is_review_only_and_no_sql_table(tmp_path):
    result = docx(tmp_path, table(row(cell('错误覆盖', '<w:vMerge/>'), cell('10'), cell('20'))))
    semantic = result['metadata']['table_semantics'][0]
    assert semantic['requires_review'] and not result['tables']
    assert any('起点' in reason for reason in semantic['review_reasons'])


def test_inconsistent_vertical_span_never_becomes_a_confirmed_grid(tmp_path):
    result = docx(tmp_path, table(
        row(cell('地区', '<w:gridSpan w:val="2"/><w:vMerge w:val="restart"/>'), cell('值'), header=True),
        row(cell('重复源文', '<w:vMerge/>'), cell('另一区域'), cell('12'))))
    semantic = result['metadata']['table_semantics'][0]
    assert not result['tables'] and semantic['requires_review']
    assert any('列跨度不一致' in reason for reason in semantic['review_reasons'])
    covered = next(c for c in semantic['cells'] if c['cell_id'] == 'R2C1')
    assert covered['text'] is None and covered['source_text'] == '重复源文'


def test_conflicting_table_units_remain_ambiguous():
    semantic = ts.build_table_semantics({'table_name': 'T', 'row_count': 2, 'column_count': 2,
        'cells': [{'row': r, 'col': c, 'text': value} for r, row_values in enumerate([['地区', '值'], ['甲', '12']], 1)
                  for c, value in enumerate(row_values, 1)]}, format='docx', explicit_header_rows=[1],
        before=[{'text': '单位：万元', 'evidence': {'page': 1, 'start': 0, 'end': 5}},
                {'text': '单位：人', 'evidence': {'page': 1, 'start': 6, 'end': 10}}])
    assert semantic['requires_review'] and {unit['scope'] for unit in semantic['units']} == {'ambiguous'}


def test_numbered_footnote_links_only_matching_header_marker_and_keeps_exact_source():
    region = {'table_name': 'VCF', 'row_count': 3, 'column_count': 3, 'cells': [
        {'row': 1, 'col': 1, 'text': 'Recovery\n(percent)', 'rowspan': 2},
        {'row': 1, 'col': 2, 'text': 'VCF', 'colspan': 2},
        {'row': 1, 'col': 3, 'text': None, 'covered_by': [1, 2]},
        {'row': 2, 'col': 1, 'text': None, 'covered_by': [1, 1]},
        {'row': 2, 'col': 2, 'text': 'Avg.1'}, {'row': 2, 'col': 3, 'text': 'Max.'},
        {'row': 3, 'col': 1, 'text': '70'}, {'row': 3, 'col': 2, 'text': '1.73'},
        {'row': 3, 'col': 3, 'text': '3.33'}]}
    note = {'text': '1 Ideal, theoretical average', 'evidence': {'page': 84, 'start': 787, 'end': 815}}
    unrelated = {'text': '2 The next section describes a different operation.', 'evidence': {'page': 84, 'start': 816, 'end': 867}}
    semantic = ts.build_table_semantics(region, format='pdf', source_position={'page': 84}, after=[note, unrelated])
    assert [p['text'] for p in semantic['columns'][1]['header_path']] == ['VCF', 'Avg.1']
    assert semantic['footnotes'] == [{**note, 'side': 'after', 'marker': '1', 'referenced_columns': [2],
                                    'role_basis': 'header_marker_and_adjacent_note_candidate', 'requires_review': True}]
    assert [(u['value'], u['scope'], u['column']) for u in semantic['units']] == [('percent', 'column', 1)]
    assert [c['text'] for c in semantic['cells'] if c['row'] == 3] == ['70', '1.73', '3.33']
    # Without a printed header marker, numbered prose must remain ordinary context.
    region['cells'][4]['text'] = 'Avg.'
    plain = ts.build_table_semantics(region, format='pdf', source_position={'page': 84}, after=[note, unrelated])
    assert not plain['footnotes'] and plain['context'][0]['text'] == note['text']


def test_xlsx_merged_paths_source_row_mapping_blank_and_zero(tmp_path):
    path = tmp_path / 'merged.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '收入'
    sheet.append(['地区', '收入', None]); sheet.merge_cells('A1:A2'); sheet.merge_cells('B1:C1')
    sheet.append([None, '本期（万元）', '同期（万元）'])
    sheet.append(['甲', 0, None]); sheet.append([None, None, None]); sheet.append(['乙', 10, 20])
    sheet.append(['注：初步统计', None, None]); book.save(path)
    result = parsers.parse(path, 'xlsx', tmp_path / 'sheet.sqlite', lambda *args: None)
    semantic = result['metadata']['table_semantics'][0]
    cells = {c['cell_id']: c for c in semantic['cells']}
    assert cells['R1C3']['covered_by'] == [1, 2] and cells['R2C1']['covered_by'] == [1, 1]
    assert cells['R3C2']['text'] == 0 and cells['R3C3']['text'] is None
    assert semantic['header_basis'] == 'unknown' and semantic['structure_unknown']
    assert all(not column['header_path'] for column in semantic['columns'])
    assert semantic['sql_row_mapping']['runs'] == [[1, 2, 2], [3, 5, 2]]
    assert semantic['requires_review'] and semantic['footnotes']
    json.dumps(result, ensure_ascii=False)
    with sqlite3.connect(result['data_path']) as database:
        assert database.execute('SELECT * FROM "收入" WHERE rowid=2').fetchone() == ('甲', 0, None)


def test_docx_caption_cuts_previous_table_unit_through_binding_context(tmp_path):
    from backend.table_context import context_for_span
    result = docx(tmp_path, table(row(cell('地区'), cell('收入'), header=True), row(cell('甲'), cell('12')), width=2)
                  + paragraph('单位：万元') + paragraph('表2 人数')
                  + table(row(cell('地区'), cell('人数'), header=True), row(cell('乙'), cell('15')), width=2))
    first, second = result['metadata']['table_semantics']
    assert first['units'][0]['scope'] == 'ambiguous'
    assert not second['units'] and second['caption']['text'] == '表2 人数'
    source_cell = next(c for c in second['cells'] if c['cell_id'] == 'R2C2')
    context = context_for_span(result, second['page'], source_cell['start'], source_cell['end'], max_chars=20000)
    projected = context['tables'][0]
    assert projected['selected_cell_ids'] == ['R2C2'] and not projected['units']


def test_xlsx_split_segments_cannot_publish_first_header_as_second_header(tmp_path):
    from backend.table_context import context_for_table
    from backend.table_bindings import check_table_bindings
    path = tmp_path / 'segments.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '双表'
    for values in [['地区', '收入（万元）'], ['甲', 12], [None, None], ['地区', '人数（人）'], ['乙', 15]]:
        sheet.append(values)
    book.save(path)
    result = parsers.parse(path, 'xlsx', tmp_path / 'segments.sqlite', lambda *args: None)
    semantic = result['metadata']['table_semantics'][0]
    assert semantic['header_basis'] == 'unknown' and semantic['structure_unknown']
    context = context_for_table(result, '双表', [3], max_chars=20000)
    projected = context['tables'][0]
    assert 'R5C2' in projected['selected_cell_ids']
    assert projected['header_basis'] == 'unknown' and all(not col['header_path'] for col in projected['columns'])
    text = '乙收入为15万元'
    verdict = check_table_bindings(text, [{'table_name': '双表', 'cell_id': 'R5C2', 'value': '15',
        'statement': text, 'column_label': '收入（万元）', 'row_label': '乙', 'unit': '万元'}],
        [{'id': 'E1', 'source_id': 'source', 'sha256': 'test', 'parser_version': parsers.PARSER_VERSION, 'table_context': context}])
    assert verdict['required'] and not verdict['valid']


def test_xlsx_unseparated_merged_headers_keep_original_paths(tmp_path):
    path = tmp_path / 'continuous.xlsx'
    book = Workbook(); sheet = book.active
    sheet.append(['地区', '收入', None]); sheet.merge_cells('A1:A2'); sheet.merge_cells('B1:C1')
    sheet.append([None, '本期（万元）', '同期（万元）']); sheet.append(['甲', 0, 12]); book.save(path)
    result = parsers.parse(path, 'xlsx', tmp_path / 'continuous.sqlite', lambda *args: None)
    semantic = result['metadata']['table_semantics'][0]
    assert semantic['header_basis'] == 'inferred'
    assert [p['text'] for p in semantic['columns'][1]['header_path']] == ['收入', '本期（万元）']


def test_unknown_text_rows_and_pure_years_do_not_become_verified_headers():
    def build(values):
        return ts.build_table_semantics({'table_name': 'T', 'row_count': len(values), 'column_count': 2,
            'cells': [{'row': r, 'col': c, 'text': value} for r, row in enumerate(values, 1) for c, value in enumerate(row, 1)]}, format='pdf')
    assert build([['甲', '牵头单位'], ['乙', '协同单位']])['header_basis'] == 'unknown'
    result = build([['地区', '年度'], ['甲', 2024], ['乙', 10]])
    assert result['header_basis'] == 'unknown' and result['requires_review']


def test_serial_number_is_location_key_not_guessed_task_name(tmp_path):
    result = docx(tmp_path, table(row(cell('序号'), cell('牵头单位'), cell('任务'), header=True),
                                row(cell('1'), cell('主管部门'), cell('建设平台'))))
    key = result['metadata']['table_semantics'][0]['row_headers'][0]
    assert key['text'] == '1' and key['column'] == 1 and key['role'] == 'row_key'


def test_metadata_budget_is_explicit_and_preserves_atomic_values(monkeypatch):
    monkeypatch.setattr(ts, 'MAX_CELLS', 3)
    result = ts.build_table_semantics({'table_name': 'T', 'row_count': 3, 'column_count': 2,
        'cells': [{'row': r, 'col': c, 'text': f'{r}-{c}'} for r in range(1, 4) for c in (1, 2)]}, format='docx')
    assert len(result['cells']) == 3 and result['truncated'] and result['requires_review']
    assert all('…' not in c['text'] for c in result['cells'])


def test_pdf_neighboring_tables_never_share_declared_units():
    first = {'table_name': 'A', 'page': 1, 'bbox': [.1, .2, .9, .4]}
    second = {'table_name': 'B', 'page': 1, 'bbox': [.1, .6, .9, .8]}
    page = {'blocks': [{'text': '单位：万元', 'bbox': [.1, .16, .4, .19], 'start': 0, 'end': 5},
                       {'text': '单位：人', 'bbox': [.1, .56, .4, .59], 'start': 6, 'end': 10}]}
    before, after = ts.page_context(first, [first, second], page)
    assert [c['text'] for c in before] == ['单位：万元']
    assert not after
    before, _ = ts.page_context(second, [first, second], page)
    assert [c['text'] for c in before] == ['单位：人']


def test_real_public_docx_grid_positions_and_textual_values(tmp_path):
    path = Path(__file__).parents[1] / 'samples/人形机器人产业方案2024_征求意见稿.docx'
    result = parsers.parse(path, 'docx', tmp_path / 'public.sqlite', lambda *args: None)
    semantic = result['metadata']['table_semantics'][0]
    assert semantic['row_count'] == 22 and semantic['column_count'] == 4
    assert semantic['header_basis'] == 'explicit'
    for c in semantic['cells']:
        if c.get('text') is not None:
            page = next(p for p in result['pages'] if p['number'] == semantic['page'])
            assert page['text'][c['start']:c['end']] == c['text']
    assert any(c['text'] == '——' for c in semantic['cells'])


def test_real_public_pdf_grid_page_two_is_bound_to_original_boxes(tmp_path, monkeypatch):
    path = Path(__file__).parents[1] / 'samples/石景山区2024年统计公报.pdf'
    monkeypatch.setattr(parsers, 'MAX_PAGES', 2)  # Deliberately bounded native-page audit.
    result = parsers.parse(path, 'pdf', tmp_path / 'public-pdf.sqlite', lambda *args: None)
    semantic = next(t for t in result['metadata']['table_semantics'] if t['page'] == 2)
    region = next(t for t in result['metadata']['table_regions'] if t['table_name'] == semantic['table_name'])
    assert semantic['row_count'] == 22 and semantic['column_count'] == 4
    assert semantic['source_position']['bbox'] == region['bbox']
    raw = {(c['row'], c['col']): c for c in region['cells']}
    assert all(c['bbox'] == raw[c['row'], c['col']]['bbox'] and c['text'] == raw[c['row'], c['col']]['text'] for c in semantic['cells'])
    assert result['metadata']['quality_summary']['processing_truncated']
